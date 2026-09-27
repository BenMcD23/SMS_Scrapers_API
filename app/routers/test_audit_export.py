"""POST /cadets/audit/export — staff-only .xlsx of the on-screen audit table."""

import io

import openpyxl
from fastapi.testclient import TestClient

from api import app
from core.security import require_staff

client = TestClient(app)

BODY = {
    "filename": "Audit - Cadet check",
    "headers": ["Surname", "CIN", "Flying Badge date"],
    "rows": [["Smith", 123, "2024-05-01"]],
}


def test_export_requires_a_token():
    assert client.post("/cadets/audit/export", json=BODY).status_code == 401


def test_export_returns_an_xlsx():
    app.dependency_overrides[require_staff] = lambda: {"email": "staff@317atc.co.uk"}
    try:
        res = client.post("/cadets/audit/export", json=BODY)
    finally:
        app.dependency_overrides.clear()
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert 'filename="Audit - Cadet check.xlsx"' in res.headers["content-disposition"]
    ws = openpyxl.load_workbook(io.BytesIO(res.content)).active
    assert [c.value for c in ws[1]] == BODY["headers"]
    assert ws["A2"].value == "Smith"
