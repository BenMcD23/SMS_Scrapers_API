"""Self-check that JI/AO signatures come from the Adult IC's saved profile
signature, and that a missing one is reported rather than silently dropped."""
import io
from types import SimpleNamespace

from docx import Document
from PIL import Image
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from database.database import Base
from database.models import User, UserSignature
from scripts.ji_ao_generator import find_signature, generate_ji


def _png() -> bytes:
    """A real image, so python-docx has something valid to embed."""
    buf = io.BytesIO()
    Image.new("RGB", (4, 2)).save(buf, "PNG")
    return buf.getvalue()


PNG = _png()


def test():
    engine = create_engine("sqlite://")
    Base.metadata.create_all(bind=engine)
    db = sessionmaker(bind=engine)()

    ben = User(google_id="1", email="b@x", first_name="Ben", last_name="McDonald")
    sam = User(google_id="2", email="s@x", first_name="Sam", last_name="McDonald")
    db.add_all([ben, sam, User(google_id="3", email="n@x", first_name="No", last_name="Sig")])
    db.flush()
    db.add_all([
        UserSignature(user_id=ben.id, image_data=PNG),
        UserSignature(user_id=sam.id, image_data=b"sam"),
    ])
    db.commit()

    assert find_signature(db, "CI Ben McDonald") == PNG
    assert find_signature(db, "Sgt Sam mcdonald") == b"sam"
    assert find_signature(db, "CI No Sig") is None
    assert find_signature(db, "N/A") is None

    event = SimpleNamespace(
        title="T", date_from=None, date_to=None, description="", location=None,
        cost=0, dress="", adult_ic="CI Ben McDonald", reference="R",
    )
    signed = Document(generate_ji(event, signature=PNG))
    assert signed.inline_shapes, "signature image should be embedded"
    unsigned = Document(generate_ji(event, signature=None))
    assert not unsigned.inline_shapes
    assert any("CI Ben McDonald RAFAC" in p.text for p in unsigned.paragraphs)


if __name__ == "__main__":
    test()
    print("ok")
