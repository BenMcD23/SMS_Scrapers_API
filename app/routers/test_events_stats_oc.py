"""Events (scraped from Bader) and JI/AO generation, squadron stats, and the
OC dashboard."""

import io
from datetime import datetime, timedelta

import pytest
from docx import Document

import routers.events as events
from database.models import (
    AllEvent,
    Cadet,
    CadetAttendance,
    CadetEvent,
    CadetQualification,
    Event317,
    Location,
    Staff,
    StaffAttendance,
    StatsSnapshot,
)

STAFF = {"Authorization": "Bearer staff"}
OC = {"Authorization": "Bearer oc"}


@pytest.fixture
def event(db):
    db.add(Location(id=1, first_line="1 Range Road", postcode="AB1 2CD"))
    db.add(Event317(id=1, title="Shooting", reference="EVT-001", adult_ic="Fg Off Smith", contact_number=7700900123,
                    date_from=datetime(2026, 5, 1, 9), date_to=datetime(2026, 5, 1, 17), location_id=1,
                    cost=12.5, dress="MTP", description="Range day"))
    db.commit()


# ── events ────────────────────────────────────────────────────────────────────

@pytest.mark.parametrize("path", ["/events", "/cadet-events", "/bans", "/generate-doc/1/fields"])
def test_event_reads_are_staff_only(api, path):
    assert api.get(path, headers=api.as_("nco")).status_code == 403


def test_events_list_serialises(api, event):
    res = api.get("/events", headers=STAFF)
    assert res.status_code == 200
    [e] = res.json()
    assert e["title"] == "Shooting" and e["reference"] == "EVT-001"


def test_cadet_events_tree(api, db):
    db.add_all([Cadet(cin=1, first_name="A", last_name="One", flight="A"),
                Cadet(cin=2, first_name="B", last_name="Two")])
    db.add_all([AllEvent(id=1, title="Camp"), AllEvent(id=2, title="Camp – Week 2", parent_id=1),
                AllEvent(id=3, title="Parade")])
    db.add_all([CadetEvent(event_id=1, cadet_id=1), CadetEvent(event_id=2, cadet_id=2)])
    db.commit()
    tree = {e["id"]: e for e in api.get("/cadet-events", headers=STAFF).json()}
    assert set(tree) == {1, 3}  # sub-apps hang off their parent, not the top level
    camp = tree[1]
    assert camp["cadet_count"] == 1 and camp["cadets"][0]["cin"] == 1
    assert camp["sub_apps"] == [{"id": 2, "title": "Camp – Week 2", "cadet_count": 1,
                                 "cadets": [{"cin": 2, "first_name": "B", "last_name": "Two",
                                             "rank": None, "flight": None}]}]
    assert tree[3]["cadets"] == [] and tree[3]["sub_apps"] == []


def test_bans_list_only_banned_cadets(api, db):
    db.add_all([Cadet(cin=1, first_name="A", last_name="One", banned=True),
                Cadet(cin=2, first_name="B", last_name="Two")])
    db.add(AllEvent(id=1, title="Camp"))
    db.add(CadetEvent(event_id=1, cadet_id=1))
    db.commit()
    assert api.get("/bans", headers=STAFF).json() == [
        {"cin": 1, "first_name": "A", "last_name": "One", "rank": None,
         "events": [{"event_id": 1, "event_title": "Camp"}]}]


def test_doc_fields_and_signature_warning(api, event):
    body = api.get("/generate-doc/1/fields", headers=STAFF).json()
    assert set(body) == {"ji", "ao", "signature_missing"}
    assert body["signature_missing"] is True
    assert isinstance(body["ji"], dict) and body["ji"]
    assert api.get("/generate-doc/99/fields", headers=STAFF).status_code == 404


@pytest.mark.parametrize("action", ["ji", "ao"])
def test_generate_doc_produces_a_word_file(api, event, action):
    res = api.post(f"/generate-doc/1/{action}", json={"fields": {"description": "Edited text"}}, headers=STAFF)
    assert res.status_code == 200
    assert f'filename="{action.upper()}_EVT-001.docx"' in res.headers["content-disposition"]
    assert res.headers["x-signature-missing"] == "1"
    text = "\n".join(p.text for p in Document(io.BytesIO(res.content)).paragraphs)
    assert "Edited text" in text or any("Edited text" in c.text for t in Document(io.BytesIO(res.content)).tables
                                         for r in t.rows for c in r.cells)


def test_generate_doc_without_a_body_uses_defaults(api, event):
    assert api.post("/generate-doc/1/ji", headers=STAFF).status_code == 200


def test_generate_doc_with_awkward_reference(api, db, event):
    db.query(Event317).update({"reference": 'EVT "2026" – Range'})
    db.commit()
    res = api.post("/generate-doc/1/ao", json={}, headers=STAFF)
    assert res.status_code == 200
    assert res.headers["content-disposition"].startswith('attachment; filename="AO_EVT 2026  Range.docx"')


def test_generate_doc_errors(api, event, monkeypatch):
    assert api.post("/generate-doc/1/memo", json={}, headers=STAFF).status_code == 400
    assert api.post("/generate-doc/99/ji", json={}, headers=STAFF).status_code == 404

    def broken(*a, **k):
        raise OSError("template missing")

    monkeypatch.setattr(events, "generate_ji", broken)
    res = api.post("/generate-doc/1/ji", json={}, headers=STAFF)
    assert res.status_code == 500 and res.json()["detail"] == "Failed to generate document"


def test_ai_description(api, event, monkeypatch):
    monkeypatch.setattr(events, "generate_ji_description_ai", lambda e: f"JI for {e.title}")
    monkeypatch.setattr(events, "generate_ao_description_ai", lambda e: f"AO for {e.title}")
    assert api.post("/generate-doc/1/ji/ai-description", headers=STAFF).json() == {"description": "JI for Shooting"}
    assert api.post("/generate-doc/1/ao/ai-description", headers=STAFF).json() == {"description": "AO for Shooting"}
    assert api.post("/generate-doc/1/memo/ai-description", headers=STAFF).status_code == 400
    assert api.post("/generate-doc/9/ji/ai-description", headers=STAFF).status_code == 404

    def down(e):
        raise RuntimeError("quota")

    monkeypatch.setattr(events, "generate_ji_description_ai", down)
    res = api.post("/generate-doc/1/ji/ai-description", headers=STAFF)
    assert res.status_code == 502 and "write it yourself" in res.json()["detail"]


# ── stats ─────────────────────────────────────────────────────────────────────

@pytest.fixture
def cadets(db):
    today = datetime.now()
    db.add_all([
        Cadet(cin=1, first_name="A", last_name="One", flight="A", rank="Cpl", classification="Leading Cadet",
              date_of_birth=datetime(today.year - 15, today.month, today.day)),
        # Birthday tomorrow: still 13 today.
        Cadet(cin=2, first_name="B", last_name="Two", flight=None,
              date_of_birth=datetime(today.year - 14, 1, 1) if (today.month, today.day) == (12, 31)
              else (datetime(today.year - 14, today.month, today.day) + timedelta(days=1))),
        Cadet(cin=3, first_name="C", last_name="Three", flight="A", classification="Junior Cadet"),
    ])
    db.add_all([
        CadetQualification(cadet_id=1, qual_type="Blue Leadership", status="true"),
        CadetQualification(cadet_id=1, qual_type="Bronze Leadership", status="true"),
    ])
    db.commit()


def test_stats_breakdown(api, cadets):
    body = api.get("/stats/current", headers=api.as_("nco")).json()
    assert body["total_cadets"] == 3
    assert body["by_flight"] == {"A": 2, "Unknown": 1}
    assert body["by_rank"] == {"Cpl": 1, "Unknown": 2}
    assert body["by_classification"] == {"Leading Cadet": 1, "Junior Cadet": 2}
    assert body["by_age"] == {"15": 1, "13": 1}
    # Highest level held counts once.
    assert body["badges"]["leadership"] == {"Bronze": 1, "None": 2}
    assert "flying_badge" in body["badges"]
    assert body["non_junior"]["total_cadets"] == 1
    assert body["non_junior"]["badges"]["leadership"] == {"Bronze": 1}


def test_stats_empty_squadron(api):
    body = api.get("/stats/current", headers=STAFF).json()
    assert body["total_cadets"] == 0 and body["by_flight"] == {}
    assert all(v == {} for v in body["badges"].values())


def test_snapshot_and_history(api, db, cadets):
    assert api.post("/stats/snapshot", headers=api.as_("nco")).status_code == 403
    res = api.post("/stats/snapshot", headers=STAFF).json()
    assert res["status"] == "ok"
    db.add(StatsSnapshot(captured_at=datetime.now() - timedelta(days=400), data={"old": True}))
    db.commit()
    history = api.get("/stats/history", headers=STAFF).json()
    assert len(history) == 1 and history[0]["data"]["total_cadets"] == 3


def test_stats_are_cached(api, db, cadets):
    api.get("/stats/current", headers=STAFF)
    db.add(Cadet(cin=9, first_name="Z", last_name="Z"))
    db.commit()
    assert api.get("/stats/current", headers=STAFF).json()["total_cadets"] == 3


# ── OC dashboard ──────────────────────────────────────────────────────────────

def test_oc_dashboard_is_oc_only(api):
    assert api.get("/oc/dashboard", headers=STAFF).status_code == 403
    assert api.get("/oc/dashboard", headers=api.as_("owner")).status_code == 403
    assert api.get("/oc/dashboard", headers=OC).status_code == 200


def test_oc_dashboard_empty(api):
    body = api.get("/oc/dashboard", headers=OC).json()
    assert body["attendance_trend"] == [] and body["expiring_quals"] == []
    assert body["qual_summary"] == {"total": 0, "expired": 0, "expiring_30": 0, "expiring_90": 0}
    assert body["strength"]["total_staff"] == 0


def test_oc_dashboard_numbers(api, db, cadets):
    now = datetime.now()
    today = datetime(now.year, now.month, now.day)
    db.add_all([
        CadetQualification(cadet_id=1, qual_type="First Aid", status="true", date_expires=today - timedelta(days=1)),
        CadetQualification(cadet_id=1, qual_type="Radio", status="true", date_expires=today + timedelta(days=10)),
        CadetQualification(cadet_id=3, qual_type="Swim", status="true", date_expires=today + timedelta(days=60)),
        CadetQualification(cadet_id=3, qual_type="DofE", status="true", date_expires=today + timedelta(days=400)),
        Staff(cin=50, first_name="S", last_name="Taff", rank="Sgt", attendance={"2026-01": 3}),
    ])
    nights = [today - timedelta(days=7 * i) for i in range(14)]
    for night in nights:
        db.add(CadetAttendance(cadet_id=1, date=night, status="Present", register_type="Parade Night"))
        db.add(CadetAttendance(cadet_id=3, date=night, status="Authorised Absence", register_type="Parade Night"))
    db.add(CadetAttendance(cadet_id=3, date=today, status="Present", register_type="Camp"))
    db.add(StaffAttendance(staff_id=50, date=nights[0], status="Present", register_type="Parade Night"))
    db.commit()

    body = api.get("/oc/dashboard", headers=OC).json()
    assert body["strength"]["total_staff"] == 1
    assert body["staff_attendance"] == [{"cin": 50, "name": "S Taff", "rank": "Sgt", "attendance": {"2026-01": 3}}]

    trend = body["attendance_trend"]
    assert len(trend) == 12
    assert trend[-1]["date"] == today.date().isoformat() and trend[0]["date"] < trend[-1]["date"]
    assert trend[-1]["cadets"] == {"present": 1, "authorised": 1, "absent": 0}
    assert trend[-1]["staff"] == {"present": 1, "authorised": 0, "absent": 0}
    assert trend[0]["staff"] == {"present": 0, "authorised": 0, "absent": 0}

    summary = body["qual_summary"]
    # The two leadership quals from the fixture have no expiry.
    assert (summary["total"], summary["expired"], summary["expiring_30"]) == (6, 1, 1)
    assert [q["qual_type"] for q in body["expiring_quals"]][0] == "Radio"
    assert body["expiring_quals"][0]["days_left"] in (9, 10)

    juniors_out = api.get("/oc/dashboard?exclude_juniors=true", headers=OC).json()
    assert juniors_out["badge_coverage"]["cadets"] == 1
    assert all(q["cadet_name"] == "A One" for q in juniors_out["expiring_quals"])
    assert juniors_out["qual_summary"]["total"] == 4
