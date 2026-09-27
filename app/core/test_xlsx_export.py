"""table_to_xlsx: header styling, date conversion, and formula-injection safety."""

import io
from datetime import datetime

import openpyxl

from core.xlsx_export import table_to_xlsx


def _sheet(data: bytes):
    return openpyxl.load_workbook(io.BytesIO(data)).active


def test_writes_headers_and_rows():
    ws = _sheet(table_to_xlsx(["Name", "CIN"], [["Smith", 123], ["Jones", None]], "Audit"))
    assert ws.title == "Audit"
    assert [c.value for c in ws[1]] == ["Name", "CIN"]
    assert ws[1][0].font.bold
    assert [c.value for c in ws[2]] == ["Smith", 123]
    assert ws["B3"].value is None
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:B3"


def test_iso_dates_become_real_dates():
    ws = _sheet(table_to_xlsx(["Awarded"], [["2024-05-01"], ["not a date"], ["2024-13-45"]]))
    assert ws["A2"].value == datetime(2024, 5, 1)
    assert ws["A2"].number_format == "DD/MM/YYYY"
    assert ws["A3"].value == "not a date"
    assert ws["A4"].value == "2024-13-45"


def test_formula_like_text_stays_literal():
    ws = _sheet(table_to_xlsx(["Note"], [["=HYPERLINK(\"http://x\")"]]))
    assert ws["A2"].data_type == "s"
    assert ws["A2"].value == "=HYPERLINK(\"http://x\")"


def test_sheet_title_is_sanitised():
    ws = _sheet(table_to_xlsx(["A"], [], "a/b:c*" + "x" * 40))
    assert ws.title == ("abc" + "x" * 40)[:31]
