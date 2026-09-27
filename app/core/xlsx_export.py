"""Turn a plain table (headers + rows) into a formatted .xlsx workbook.

Shared so any "export this list to Excel" feature gets the same presentation
(bold frozen header, autofilter, sized columns, real date cells) without each
router hand-rolling openpyxl styling.
"""

import io
import re
from datetime import date

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

_ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_MAX_COL_WIDTH = 60


def _cell_value(value):
    # ISO dates become real Excel dates so the sheet sorts and filters them
    # chronologically rather than as text.
    if isinstance(value, str) and _ISO_DATE.match(value):
        try:
            return date.fromisoformat(value)
        except ValueError:
            return value
    return value


def table_to_xlsx(headers: list[str], rows: list[list], sheet_title: str = "Sheet1") -> bytes:
    """A single-sheet workbook of ``rows`` under ``headers``, as bytes."""
    wb = openpyxl.Workbook()
    ws = wb.active
    # Excel caps sheet names at 31 chars and forbids a handful of characters.
    ws.title = re.sub(r"[\[\]:*?/\\]", "", sheet_title)[:31] or "Sheet1"

    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E5E7EB")

    for row in rows:
        ws.append([_cell_value(v) for v in row])

    widths = [len(str(h)) for h in headers]
    for row in ws.iter_rows(min_row=2):
        for i, cell in enumerate(row):
            if isinstance(cell.value, date):
                cell.number_format = "DD/MM/YYYY"
                length = 10
            else:
                # openpyxl treats any string starting with "=" as a formula;
                # exported names/notes must stay literal text, never execute.
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"
                length = max((len(line) for line in str(cell.value or "").splitlines()), default=0)
            if i < len(widths):
                widths[i] = max(widths[i], length)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(w + 2, _MAX_COL_WIDTH)

    ws.freeze_panes = "A2"
    if headers:
        ws.auto_filter.ref = ws.dimensions

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
