"""Form and sheet generators: the Logs Form 202 demand sheet, the HTD money
logic, and the AI-text parsers. The Logs Form checks are the module's own
__main__ self-check, moved here so CI actually runs them."""

import io
from datetime import date, timedelta

import openpyxl
import pytest

from core.catalogue import UNIFORM_SIZES
from form_generators.HTD_gen import compute_htd
from form_generators.logs_form_gen import TEMPLATE_PATH, _norm, build_description, generate_logs_form
from scripts.nco_appraisal_ai import _clean, _split_sections, short_name


def _live_demands(xlsx: bytes) -> dict:
    ws = openpyxl.load_workbook(io.BytesIO(xlsx))["Live Demands"]
    return {ws.cell(r, 4).value: (ws.cell(r, 5).value, ws.cell(r, 7).value)
            for r in range(2, ws.max_row + 1) if ws.cell(r, 4).value not in (None, "")}


def test_every_catalogue_size_has_a_template_demand_line():
    ws = openpyxl.load_workbook(TEMPLATE_PATH)["Live Demands"]
    rows = {_norm(ws.cell(r, 4).value) for r in range(2, ws.max_row + 1) if ws.cell(r, 4).value not in (None, "")}
    catalogue = {**UNIFORM_SIZES, "Belt": [""]}
    missing = [(t, s) for t, sizes in catalogue.items() for s in sizes
               if build_description(t, s) is not None and _norm(build_description(t, s)) not in rows]
    # Known template typo: the template says 80/70/99 where the catalogue has 80/75/99.
    assert missing == [("Slacks", "80/75/99")]


@pytest.fixture(scope="module")
def logs_form():
    entries = [("Beret", "56"), ("Beret", "56"), ("Tie", "Standard"), ("Slacks", "80/75/99"), ("Brassard", "")]
    roll = [("Cdt", "A Cadet", "Initial Issue"), ("", "B Adult", "")]
    return generate_logs_form(entries, roll)


def test_logs_form_demands_only_what_was_ordered(logs_form):
    rows = _live_demands(logs_form)
    assert {d: q for d, (q, _) in rows.items()} == {
        "Beret RAF Size 56": 2,
        "Necktie Black (Unisex Item) Standard": 1,
        "Trousers Woman's RAF 80/75/99": 1,  # no template row, so appended
    }
    rdds = {v.date() if hasattr(v, "date") else v for _, v in rows.values()}
    assert rdds == {date.today() + timedelta(days=21)}


def test_logs_form_nominal_roll(logs_form):
    nr = openpyxl.load_workbook(io.BytesIO(logs_form))["Cadet Nominal Roll"]
    assert [nr.cell(3, c).value for c in range(1, 5)] == [1, "Cdt", "A Cadet", "Initial Issue"]
    assert [nr.cell(4, c).value for c in range(1, 5)] == [2, None, "B Adult", None]  # openpyxl stores "" as empty
    # The template's placeholder rows below are cleared.
    assert all(nr.cell(r, 3).value is None for r in range(5, nr.max_row + 1))


def test_build_description():
    assert build_description("Beret", "55") == "Beret RAF Size 55"
    assert build_description("Belt", "anything") == "Belt Waist RAF (Unisex Item) Size 64-114cm"
    assert build_description("Brassard", "") is None


@pytest.mark.parametrize("a,b", [("Female115/42 ", "female 115/42"), ("Woman’s", "Woman's")])
def test_norm_tolerates_template_typos(a, b):
    assert _norm(a) == _norm(b)


# ── HTD money ─────────────────────────────────────────────────────────────────

@pytest.mark.parametrize("distance,journeys,car,total", [
    (10, [8, 6], 2.68, 37.52),
    (0, [5], 0.0, 0.0),
    (12.5, [], 3.34, 0.0),
    (100, [1, 1, 1, 1, 1, 1], 26.75, 160.5),
])
def test_compute_htd(distance, journeys, car, total):
    out = compute_htd(distance, journeys)
    assert out["car_cost"] == car and out["total_a"] == car
    assert out["amounts"] == out["totals"] == [round(j * car, 2) for j in journeys]
    assert out["total_claimed"] == total


# ── NCO appraisal AI output parsing ───────────────────────────────────────────

def test_split_sections_in_any_order_and_missing_ones_blank():
    out = ("Sure! ===STRENGTHS===\n- **Calm** under pressure\n\n\n* Reliable\n"
           "===GENERAL_OBSERVATIONS===\nA solid NCO. === not a marker\n===TARGETS===\n• Delegate")
    sections = _split_sections(out)
    assert sections["strengths"] == "Calm under pressure\n\nReliable"
    assert sections["general_observations"] == "A solid NCO. === not a marker"
    assert sections["targets"] == "Delegate"
    assert sections["weaknesses"] == "" and sections["effectiveness_in_role"] == ""


def test_split_sections_with_no_markers():
    assert set(_split_sections("I can't help with that").values()) == {""}


def test_clean_collapses_blank_runs():
    assert _clean("\n\n- a\n\n\n\n- b\n\n") == "a\n\nb"


@pytest.mark.parametrize("name,short", [
    ("Corporal Isabella Wiggett", "Cpl Wiggett"),
    ("cpl Isabella Wiggett", "Cpl Wiggett"),
    ("Flight Sergeant Al Jones", "FS Jones"),
    ("Isabella Wiggett", "Isabella Wiggett"),
    ("", "the NCO"),
    (None, "the NCO"),
])
def test_short_name(name, short):
    assert short_name(name) == short
